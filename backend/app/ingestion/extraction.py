import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from app.ingestion.classifier import fold
from app.ingestion.financial_schemas import (
    BilanExtraction,
    FinancialTable,
    InvoiceExtraction,
    LineItem,
    QuoteExtraction,
    Source,
    SourcedValue,
)
from app.ingestion.normalized import NormalizedBlock, NormalizedDocument
from openpyxl.utils import get_column_letter

NUMBER = r"\(?-?\d[\d\s\u00a0\u202f.,]*\)?"


def parse_number(text: str) -> float | None:
    """Decimal arithmetic input with French/English grouping; never evaluate formulas."""
    text = text.strip().replace("\u2212", "-")
    if "[formula:" in text or not re.search(r"\d", text):
        return None
    match = re.search(NUMBER, text)
    if not match:
        return None
    token = re.sub(r"\s", "", match.group())
    negative = token.startswith("(")
    token = token.strip("()")
    if "," in token and "." in token:
        decimal = "," if token.rfind(",") > token.rfind(".") else "."
        token = token.replace("." if decimal == "," else ",", "").replace(decimal, ".")
    elif "," in token:
        parts = token.split(",")
        token = "".join(parts) if len(parts[-1]) == 3 else "".join(parts[:-1]) + "." + parts[-1]
    elif token.count(".") > 1:
        token = token.replace(".", "")
    try:
        result = Decimal(token)
        if negative:
            result = -result
        return float(result) if result.is_finite() else None
    except InvalidOperation:
        return None


@dataclass
class EvidenceLine:
    text: str
    block: NormalizedBlock
    cell_range: str | None = None


def evidence_lines(doc: NormalizedDocument) -> list[EvidenceLine]:
    lines = []
    for block in doc.blocks:
        rows = block.metadata.get("rows")
        if rows:
            for i, row in enumerate(rows):
                location = block.cell_range
                if block.sheet_name:
                    row_index = block.metadata.get("row_start", 1) + i
                    location = f"A{row_index}:{get_column_letter(max(1, len(row)))}{row_index}"
                lines.append(EvidenceLine(" | ".join(map(str, row)), block, location))
        else:
            lines.extend(EvidenceLine(line, block) for line in block.text.splitlines() if line.strip())
    return lines


def source_for(doc: NormalizedDocument, line: EvidenceLine) -> Source:
    block = line.block
    return Source(
        document_id=doc.document_id,
        filename=doc.filename,
        block_id=block.block_id,
        page=block.page,
        table_id=block.table_id,
        sheet=block.sheet_name,
        cell_range=line.cell_range or block.cell_range,
        section=block.section,
        slide_number=block.slide_number,
        excerpt=line.text,
    )


class StructuredExtractor:
    schemas = {"invoice": InvoiceExtraction, "devis": QuoteExtraction, "bilan": BilanExtraction}

    def extract(self, doc: NormalizedDocument) -> dict:
        schema = self.schemas.get(doc.document_type)
        if not schema:
            return {"warnings": [], "financial_tables": []}
        result = schema()
        lines = evidence_lines(doc)

        def find(labels: str, numeric: bool = False) -> SourcedValue | None:
            for line in lines:
                text = fold(line.text)
                match = re.match(r"^\s*(?:" + labels + r")\s*(?:\s*[:|#]\s*|\s+)(.+?)\s*\|?\s*$", text)
                if not match:
                    continue
                # Accent folding preserves character positions for ordinary Latin text.
                original = line.text[match.start(1) : match.end(1)].strip(" |:")
                if numeric and not re.match(r"^(?:(?:EUR|USD|GBP|TND|[$€£])\s*)?[\d(+-]", original, re.I):
                    continue
                value = parse_number(original) if numeric else original
                if value is not None and value != "":
                    return SourcedValue(value=value, source=source_for(doc, line))
            return None

        for line in lines:
            currency_match = re.search(r"\b(EUR|USD|GBP|TND|CAD|CHF|MAD|DZD)\b|[€£]", line.text, re.I)
            if currency_match:
                symbol = currency_match.group().upper()
                result.currency = SourcedValue(
                    value={"€": "EUR", "£": "GBP"}.get(symbol, symbol), source=source_for(doc, line)
                )
                break
        if isinstance(result, BilanExtraction):
            labels = {
                "company_name": "societe|company(?: name)?|entreprise",
                "reporting_period": "exercice|reporting period|fiscal year|annee",
                "units": "unites|units",
                "assets": "total actif|total assets|actif|assets",
                "current_assets": "actif courant|current assets",
                "non_current_assets": "actif non courant|non.current assets",
                "liabilities": "total passif|total liabilities|passif|liabilities",
                "current_liabilities": "passif courant|current liabilities",
                "non_current_liabilities": "passif non courant|non.current liabilities",
                "equity": "capitaux propres|equity|total equity",
                "revenue": "chiffre d.affaires|revenue|net revenue",
                "operating_income": "resultat d.exploitation|operating income",
                "net_income": "resultat net|net income|net profit",
                "cash": "tresorerie|cash(?: and cash equivalents)?",
            }
            for field, pattern in labels.items():
                setattr(
                    result,
                    field,
                    find(pattern, numeric=field not in ("company_name", "reporting_period", "units")),
                )
            result.financial_tables = [
                FinancialTable(
                    rows=[[str(c) for c in row] for row in b.metadata["rows"]],
                    source=source_for(doc, EvidenceLine(b.text, b)),
                )
                for b in doc.blocks
                if b.metadata.get("rows")
            ]
        else:
            for field, labels in {
                "subtotal": "total ht|subtotal|sub.total|sous.total|montant ht",
                "tax_amount": r"montant tva|tax amount|total tax|tva(?: \(?\d+\s*%\)?)?",
                "total": "montant ttc|total ttc|grand total|amount due|total due|total",
            }.items():
                setattr(result, field, find(labels, numeric=True))
            result.payment_terms = find("conditions de paiement|payment terms")
            for party, prefix in (
                (result.supplier, "fournisseur|supplier|bill from"),
                (result.customer, "client|customer|bill to"),
            ):
                party.name = find(prefix)
                for field, label in {
                    "address": "adresse|address",
                    "tax_id": "tax id|vat id|numero tva|matricule fiscal",
                    "registration_number": "registration number|siret",
                    "email": "email|courriel",
                    "phone": "phone|telephone",
                }.items():
                    party_value = find(f"(?:{prefix}) (?:{label})")
                    if party is result.supplier and party_value is None:
                        party_value = find(label)
                    setattr(party, field, party_value)
            result.line_items = self._line_items(doc)
            if isinstance(result, InvoiceExtraction):
                result.invoice_number = find(
                    r"invoice number|invoice no\.?|numero de facture|facture n[o°.]?|facture|invoice"
                )
                result.invoice_date = find("date de facture|invoice date|date")
                result.due_date = find("date d.echeance|echeance|due date")
                result.iban = find("iban")
            else:
                result.quote_number = find(
                    "quote number|quotation number|numero de devis|devis n[o°.]?|devis|quote|quotation"
                )
                result.quote_date = find("date du devis|quote date|quotation date|date")
                result.valid_until = find("valable jusqu.au|valid until")
                result.validity = find("validite(?: du devis)?|validity")
                result.notes = find("notes|remarques")
            self._validate_arithmetic(result)
        return result.model_dump(mode="json")

    def _line_items(self, doc: NormalizedDocument) -> list[LineItem]:
        items = []
        labels = {
            "description": r"description|designation|item",
            "quantity": r"quantity|quantite|qty|qte",
            "unit": r"unit|unite",
            "unit_price": r"unit price|prix unitaire|pu(?: ht)?",
            "tax_rate": r"tax rate|tva|vat",
            "amount": r"amount|montant(?: ht)?|line total|total",
        }
        for block in doc.blocks:
            rows = block.metadata.get("rows", [])
            for header_index, header in enumerate(rows):
                columns = {
                    field: i
                    for field, label in labels.items()
                    for i, cell in enumerate(header)
                    if re.fullmatch(label, fold(str(cell)).strip())
                }
                if "description" not in columns or "amount" not in columns:
                    continue
                for offset, row in enumerate(rows[header_index + 1 :], header_index + 1):
                    if not row or not str(row[columns["description"]]).strip():
                        continue
                    values = {}
                    for field, col in columns.items():
                        if col >= len(row):
                            continue
                        value = (
                            str(row[col]) if field in ("description", "unit") else parse_number(str(row[col]))
                        )
                        if value is not None and value != "":
                            cell_range = (
                                f"{get_column_letter(col + 1)}{block.metadata.get('row_start', 1) + offset}"
                                if block.sheet_name
                                else None
                            )
                            values[field] = SourcedValue(
                                value=value,
                                source=source_for(
                                    doc, EvidenceLine(" | ".join(map(str, row)), block, cell_range)
                                ),
                            )
                    if values.get("amount"):
                        items.append(LineItem(**values))
                break
        return items

    @staticmethod
    def _validate_arithmetic(result):
        if all(getattr(result, name) is not None for name in ("subtotal", "tax_amount", "total")):
            subtotal, tax, total = (
                Decimal(str(getattr(result, name).value)) for name in ("subtotal", "tax_amount", "total")
            )
            if abs(subtotal + tax - total) > Decimal("0.02"):
                result.warnings.append(
                    "subtotal + tax_amount does not equal total (tolerance 0.02). Values were preserved."
                )
        for index, item in enumerate(result.line_items, 1):
            if item.quantity and item.unit_price and item.amount:
                calculated = Decimal(str(item.quantity.value)) * Decimal(str(item.unit_price.value))
                if abs(calculated - Decimal(str(item.amount.value))) > Decimal("0.02"):
                    result.warnings.append(
                        f"Line {index}: quantity × unit_price does not equal amount. Values were preserved."
                    )
