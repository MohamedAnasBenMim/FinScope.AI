import re
from pathlib import Path

import pdfplumber
import pypdfium2
from app.config import Settings
from app.errors import ApplicationError, ParsingFailed
from app.ingestion.normalized import NormalizedBlock, NormalizedDocument, table_markdown
from app.ingestion.ocr import BaseOCRProvider, TesseractOCRProvider
from app.ingestion.spreadsheets import parse_spreadsheet
from app.ingestion.validation import decode_text
from bs4 import BeautifulSoup
from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageSequence
from pptx import Presentation


class DocumentParser:
    def __init__(self, config: Settings, ocr: BaseOCRProvider | None = None):
        self.config = config
        self.ocr = ocr or TesseractOCRProvider(config)

    def parse(
        self, path: Path, document_id: str, filename: str, mime: str, extension: str
    ) -> NormalizedDocument:
        doc = NormalizedDocument(
            document_id=document_id,
            filename=filename,
            mime_type=mime,
            extension=extension,
            metadata={
                "parser": self.config.PARSER_PROVIDER,
                "ocr_used": False,
                "ocr_engine": None,
                "ocr_confidence": None,
            },
        )
        try:
            if self.config.PARSER_PROVIDER == "docling" and extension in (".pdf", ".docx", ".pptx", ".html"):
                from app.ingestion.docling_parser import convert_docling

                return convert_docling(path, doc, self.config)
            if extension == ".pdf":
                self._pdf(path, doc)
            elif mime.startswith("image/"):
                with Image.open(path) as image:
                    frame_count = getattr(image, "n_frames", 1)
                    doc.metadata["image_count"] = frame_count
                    for index, frame in enumerate(ImageSequence.Iterator(image), 1):
                        self._ocr(frame.copy(), doc, index if frame_count > 1 else None)
            elif extension in (".xlsx", ".xls", ".csv"):
                doc.blocks = parse_spreadsheet(path, extension, self.config)
                doc.metadata["workbook"] = filename
            elif extension == ".docx":
                word = WordDocument(path)
                heading = None
                for item in word.iter_inner_content():
                    if isinstance(item, Paragraph) and item.text.strip():
                        is_heading = item.style is not None and item.style.name.startswith(
                            ("Heading", "Title")
                        )
                        if is_heading:
                            heading = item.text
                        doc.blocks.append(
                            NormalizedBlock(
                                type="heading" if is_heading else "text",
                                text=item.text,
                                heading=heading,
                                section=heading,
                            )
                        )
                    elif isinstance(item, Table):
                        rows = [[cell.text for cell in row.cells] for row in item.rows]
                        doc.blocks.append(
                            NormalizedBlock(
                                type="table",
                                text=table_markdown(rows),
                                heading=heading,
                                section=heading,
                                table_id=f"table-{len(doc.blocks)}",
                                metadata={"rows": rows, "header_rows": 1},
                            )
                        )
            elif extension == ".pptx":
                presentation = Presentation(path)
                if len(presentation.slides) > self.config.MAX_DOCUMENT_PAGES:
                    raise ParsingFailed("Presentation exceeds MAX_DOCUMENT_PAGES.")
                doc.metadata["slide_count"] = len(presentation.slides)
                for index, slide in enumerate(presentation.slides, 1):
                    heading = slide.shapes.title.text if slide.shapes.title is not None else None
                    for shape in slide.shapes:
                        if shape.has_table:
                            rows = [[c.text for c in r.cells] for r in shape.table.rows]
                            doc.blocks.append(
                                NormalizedBlock(
                                    type="table",
                                    text=table_markdown(rows),
                                    slide_number=index,
                                    heading=heading,
                                    table_id=f"slide-{index}-{shape.shape_id}",
                                    metadata={"rows": rows, "header_rows": 1},
                                )
                            )
                        elif shape.has_text_frame and shape.text.strip():
                            doc.blocks.append(
                                NormalizedBlock(
                                    text=shape.text, slide_number=index, heading=heading, section=heading
                                )
                            )
            elif extension in (".html", ".htm"):
                soup = BeautifulSoup(decode_text(path.read_bytes()), "html.parser")
                for tag in soup(["script", "style", "iframe", "object", "noscript"]):
                    tag.decompose()
                heading = None
                for tag in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "table"]):
                    if tag.find_parent("table"):
                        continue
                    if tag.name == "table":
                        rows = [
                            [c.get_text(" ", strip=True) for c in r.find_all(["td", "th"])]
                            for r in tag.find_all("tr")
                        ]
                        doc.blocks.append(
                            NormalizedBlock(
                                type="table",
                                text=table_markdown(rows),
                                section=heading,
                                table_id=f"table-{len(doc.blocks)}",
                                metadata={"rows": rows, "header_rows": 1},
                            )
                        )
                    else:
                        text = tag.get_text(" ", strip=True)
                        is_heading = tag.name.startswith("h")
                        if is_heading:
                            heading = text
                        if text:
                            doc.blocks.append(
                                NormalizedBlock(
                                    type="heading" if is_heading else "text",
                                    text=text,
                                    section=heading,
                                    heading=heading,
                                )
                            )
                if not doc.blocks and soup.get_text(" ", strip=True):
                    doc.blocks = [NormalizedBlock(text=soup.get_text(" ", strip=True))]
            else:
                doc.blocks = self._text_blocks(decode_text(path.read_bytes()))
            if not any(b.text.strip() for b in doc.blocks):
                raise ParsingFailed(
                    "No readable content found after parsing/OCR. Supply a clearer scan or a document with text."
                )
            return doc.assign_ids()
        except ApplicationError:
            raise
        except Exception as exc:
            raise ParsingFailed(
                f"Document parsing failed ({type(exc).__name__}). The file may be malformed."
            ) from exc

    def _ocr(self, image: Image.Image, doc: NormalizedDocument, page: int | None):
        result = self.ocr.recognize(image, page)
        doc.blocks.extend(result.blocks)
        doc.metadata.update(ocr_used=True, ocr_engine=result.engine)
        confidences = doc.metadata.setdefault("ocr_page_confidences", [])
        if result.confidence is not None:
            confidences.append(result.confidence)
            doc.metadata["ocr_confidence"] = sum(confidences) / len(confidences)

    def _pdf(self, path: Path, doc: NormalizedDocument):
        with pdfplumber.open(path) as pdf:
            doc.page_count = len(pdf.pages)
            for index, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if len(re.sub(r"\s", "", text)) < self.config.OCR_MIN_TEXT_CHARS:
                    with pypdfium2.PdfDocument(str(path)) as rendered:
                        rendered_page = rendered[index - 1]
                        width, height = rendered_page.get_size()
                        scale = min(2.5, (self.config.MAX_IMAGE_PIXELS / (width * height)) ** 0.5)
                        bitmap = rendered_page.render(scale=scale)
                        self._ocr(bitmap.to_pil(), doc, index)
                        bitmap.close()
                        rendered_page.close()
                    continue
                tables = page.find_tables()
                for table_index, table in enumerate(tables, 1):
                    rows = [[c or "" for c in row] for row in table.extract()]
                    doc.blocks.append(
                        NormalizedBlock(
                            type="table",
                            text=table_markdown(rows),
                            page=index,
                            table_id=f"p{index}-table-{table_index}",
                            bbox=list(table.bbox),
                            metadata={"rows": rows, "header_rows": 1, "bbox_units": "pdf_points"},
                        )
                    )
                # Reuse pdfplumber's geometry to avoid duplicating table cells in narrative chunks.
                narrative = (
                    page.filter(
                        lambda obj: (
                            not any(
                                t.bbox[0] <= obj.get("x0", -1) <= t.bbox[2]
                                and t.bbox[1] <= obj.get("top", -1) <= t.bbox[3]
                                for t in tables
                            )
                        )
                    ).extract_text()
                    or ""
                )
                for block in self._text_blocks(narrative):
                    block.page = index
                    doc.blocks.append(block)

    @staticmethod
    def _text_blocks(text: str) -> list[NormalizedBlock]:
        blocks = []
        heading = None
        lines = text.splitlines()
        index = 0
        while index < len(lines):
            line = lines[index].strip()
            if not line:
                index += 1
                continue
            if "|" in line and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3}", lines[index + 1]):
                rows = [[c.strip() for c in line.strip("|").split("|")]]
                index += 2
                while index < len(lines) and "|" in lines[index]:
                    rows.append([c.strip() for c in lines[index].strip("|").split("|")])
                    index += 1
                blocks.append(
                    NormalizedBlock(
                        type="table",
                        text=table_markdown(rows),
                        heading=heading,
                        section=heading,
                        table_id=f"table-{len(blocks)}",
                        metadata={"rows": rows, "header_rows": 1},
                    )
                )
                continue
            is_heading = line.startswith("#") or (line.isupper() and len(line) < 100)
            if is_heading:
                heading = line.lstrip("# ").strip()
            blocks.append(
                NormalizedBlock(
                    type="heading" if is_heading else "text",
                    text=line,
                    heading=heading,
                    section=heading,
                    metadata={"line_start": index + 1, "line_end": index + 1},
                )
            )
            index += 1
        return blocks
